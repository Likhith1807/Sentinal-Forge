import copy, io, sys
sys.path.insert(0, __import__("os").path.dirname(__import__("os").path.abspath(__file__)))
from lib import *
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

SRC = "D:/SENTINEL_Forge/format.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "D:/SENTINEL_Forge/SENTINEL_Forge_Final_Presentation.pptx"

prs = Presentation(SRC)
S = list(prs.slides)              # original 13 slides, indexed 0..12
BASE = S[2]                       # 'Problem Statement' slide is the skeleton for every content slide
KEEP_PREFIX = ("Slide Number", "Footer", "Picture 12", "Group 8")

def is_title(sh):
    if not sh.has_text_frame or sh.top > Inches(0.7) or sh.left < Inches(1.0): return False
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            return r.font.size is not None and r.font.size.pt == 32
    return False

def clear_body(slide):
    for sh in list(slide.shapes):
        if sh.name.startswith(KEEP_PREFIX) or is_title(sh): continue
        sh._element.getparent().remove(sh._element)

def set_title(slide, text):
    for sh in slide.shapes:
        if is_title(sh):
            runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
            sh.width = Inches(11.6)
            runs[0].text = text
            runs[0].font.name = SERIF; runs[0].font.bold = True; runs[0].font.size = Pt(32)
            for cs in runs[0]._r.rPr.findall(qn("a:cs")): runs[0]._r.rPr.remove(cs)
            etree.SubElement(runs[0]._r.rPr, qn("a:cs")).set("typeface", SERIF)
            for r in runs[1:]: r.text = ""
            return sh
    raise RuntimeError("no title")

def clone_skeleton(title):
    """New slide = exact copy of the format's content-slide chrome (logo, 32pt title, double rule, footer, number)."""
    layout = BASE.slide_layout
    s = prs.slides.add_slide(layout)
    for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
    logo_rid = None
    for sh in BASE.shapes:
        if sh.name.startswith(KEEP_PREFIX) or is_title(sh):
            el = copy.deepcopy(sh._element)
            if sh.name == "Picture 12":
                blob = sh.image.blob
                _, rid = s.part.get_or_add_image_part(io.BytesIO(blob))
                el.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
            s.shapes._spTree.append(el)
    set_title(s, title)
    return s

def prep(slide, title):
    clear_body(slide); set_title(slide, title)
    for tm in slide._element.findall(qn("p:timing")):      # animations pointed at shapes we just replaced
        slide._element.remove(tm)

# ------------------------------------------------------------------ 1. TITLE
s1 = S[0]
def set_para_text(shape, texts):
    """replace text run-by-run, keeping each paragraph's first-run formatting"""
    paras = shape.text_frame.paragraphs
    for p, t in zip(paras, texts):
        runs = p.runs
        if not runs: continue
        runs[0].text = t
        for r in runs[1:]: r.text = ""
    for p in paras[len(texts):]:
        p._p.getparent().remove(p._p)
for sh in s1.shapes:
    if not sh.has_text_frame: continue
    t = sh.text_frame.text
    if t.startswith("Name of Topic"):
        set_para_text(sh, ["SENTINEL Forge: Evidence-Grounded Compilation of Threat Reports into Tested Detection Rules"])
    elif t.startswith("B.Tech"):
        set_para_text(sh, ["B.Tech (AIDS)", "UG-II (IV-Semester)"])
    elif t.startswith("Subject Coordinator"):
        set_para_text(sh, ["Subject Coordinator", "Dr. ____________", "Professor"])
    elif t.startswith("Name of Student"):
        set_para_text(sh, ["K. Karthikeya (DL.AI.U4AID24020)", "A. Karthik (DL.AI.U4AID24104)",
                           "K. Likhith (DL.AI.U4AID24117)", "Y. Sandeep (DL.AI.U4AID24143)"])
        sh.left = Inches(8.12); sh.width = Inches(5.1)
c = chip(s1, 2.5, 2.96, 8.33, 0.44, "Course Project  ·  Text Analytics + Big Data Analytics", fill=ROSE, color=MAROON, size=18, font=SERIF, line=ROSE2)
notes(s1, "Introduce SENTINEL Forge in one sentence: it turns written threat reports into detection rules that are checked against the available logs and replay-tested. "
          "Mention that the project spans two courses: Text Analytics (NLP extraction) and Big Data Analytics (Scala compiler + Apache Spark).")

# ------------------------------------------------------------------ 2. CONTENTS
s2 = S[1]
prep(s2, "CONTENTS")
items = [("Project Idea", 3), ("Problem Statement", 4), ("Literature Review (Previous Work)", 5), ("Novelty of the Project", 6),
         ("Objectives of the Project", 7), ("Proposed Methodology", 8), ("Implementation Layout", 9), ("Tools and Technologies", 10),
         ("Experimental Setup", 11), ("Results", "12–13"), ("Challenges", 14), ("Project Timeline", 15),
         ("Conclusions", 16), ("Future Work & Scope", 17), ("References", 18)]
for i, (label, pg) in enumerate(items):
    col, row = divmod(i, 8)
    x = 0.45 + col * 6.4; y = 1.42 + row * 0.66
    box(s2, x, y, 6.1, 0.54, fill=ROSE if i % 2 == 0 else PAPER, line=None, radius=0.3)
    num_circle(s2, i + 1, x + 0.08, y + 0.06, 0.42, size=16)
    tb(s2, x + 0.62, y, 4.6, 0.54, [{"runs": [label]}], anchor="m", size=22, font=SERIF)
    tb(s2, x + 4.6, y, 1.4, 0.54, [{"runs": [f"Slide {pg}"], "align": "r"}], anchor="m", size=13, color=GREY)
notes(s2, "Walk through the agenda. The deck follows the department format, with four additions the brief requires: Project Idea, Novelty, Implementation Layout and Project Timeline, plus a dedicated Future Work slide.")

# ------------------------------------------------------------------ 3. PROJECT IDEA (new)
s = clone_skeleton("Project Idea")
lead(s, ["Turn written threat reports into ", ("tested, evidence-linked", {"bold": True, "color": MAROON}), " detection rules."])
cw = 4.01
cards = [("FaSearch", "1  ·  Understand", "NLP extracts attacker actions, entities, thresholds and event sequences from a report, keeping the exact source sentence as evidence."),
         ("FaCogs", "2  ·  Compile", "Checks that the log schema can actually observe the behaviour, then emits a typed Scala/Spark rule from a closed set of recipes."),
         ("FaClipboardCheck", "3  ·  Validate", "Replays independently labelled events, reports misses and false alarms, and asks an analyst to approve, or explains why no rule can be made.")]
for i, (ic, t, d) in enumerate(cards):
    x = 0.4 + i * (cw + 0.245)
    card(s, x, 1.95, cw, 2.25)
    icon_circle(s, ic, x + 0.22, 2.15, 0.66)
    tb(s, x + 1.0, 2.15, cw - 1.15, 0.66, [{"runs": [t]}], anchor="m", size=21, bold=True, font=SERIF, color=MAROON)
    tb(s, x + 0.22, 3.0, cw - 0.44, 1.15, [{"runs": [d]}], size=15, color=INK)
    if i < 2: chevron(s, x + cw - 0.02, 2.93, 0.29, 0.42)
box(s, 0.4, 4.42, 12.53, 2.48, fill=NAVY_T, radius=0.05)
tb(s, 0.62, 4.5, 6, 0.36, [{"runs": ["WORKED EXAMPLE  ·  repeated failed sign-ins, then a success"]}], anchor="m", size=13, bold=True, color=NAVY)
steps = [("Report sentence", "“Five or more failed sign-ins within two minutes, then a successful sign-in for the same account.”"),
         ("Typed spec", "SequenceThenTrigger · group by account_id · ≥ 5 login_failure in 120 s · trigger on login_success"),
         ("Spark rule", "Trailing-window count of failures per account; alert on the next successful login."),
         ("Replay verdict", "6 / 6 labelled scenarios correct, including exact window-boundary cases.")]
bw = 2.78; gap = (12.09 - 4 * bw) / 3
for i, (h, d) in enumerate(steps):
    x = 0.62 + i * (bw + gap)
    card(s, x, 4.95, bw, 1.8, radius=0.06)
    tb(s, x + 0.12, 5.02, bw - 0.24, 0.36, [{"runs": [h]}], anchor="m", size=16, bold=True, font=SERIF, color=MAROON)
    tb(s, x + 0.12, 5.4, bw - 0.24, 1.3, [{"runs": [d], "italic": i == 0}], size=13.5, color=INK)
    if i < 3: chevron(s, x + bw + gap / 2 - 0.13, 5.68, 0.26, 0.36, color=NAVY)
notes(s, "The idea: a threat report says what an attacker does, but a usable detection needs the right log fields and conditions. SENTINEL Forge automates that translation in three steps - understand, compile, validate - and it refuses to produce a rule when the logs cannot observe the behaviour. The worked example is behaviour B1 from our five-behaviour scope.")
S_IDEA = s

# ------------------------------------------------------------------ 4. PROBLEM STATEMENT
s = S[2]; prep(s, "Problem Statement")
lead(s, ["A threat report says ", ("what an attacker did", {"bold": True, "color": MAROON}), ", not how to detect it in logs."])
probs = [("FaFileAlt", "Manual translation gap", "Analysts read prose reports, then hand-write a rule for every SIEM. It is slow, inconsistent and hard to audit."),
         ("FaEye", "Conditions the logs cannot see", "A rule can need a field the logs never record (e.g. source_ip is often dropped). It runs cleanly yet detects nothing."),
         ("FaBug", "Plausible but wrong LLM rules", "Direct LLM generation compiles, but hard-codes account names, uses wrong enum values or floods duplicate alerts.")]
for i, (ic, t, d) in enumerate(probs):
    y = 1.95 + i * 1.7
    card(s, 0.4, y, 7.45, 1.5)
    icon_circle(s, ic, 0.6, y + 0.4, 0.7)
    tb(s, 1.5, y + 0.1, 6.2, 0.42, [{"runs": [t]}], anchor="m", size=20, bold=True, font=SERIF, color=MAROON)
    tb(s, 1.5, y + 0.55, 6.2, 0.9, [{"runs": [d]}], size=15)
tb(s, 8.1, 1.9, 4.83, 0.32, [{"runs": ["MEASURED ON HELD-OUT REPORTS"]}], anchor="m", size=12.5, bold=True, color=GREY)
stats = [("1/5", RED, RED_T, "direct-LLM rules fully correct"), ("0/5", RED, RED_T, "schema-constrained rules correct on every scenario"),
         ("17/17", GREEN, GREEN_T, "replay scenarios passed by SENTINEL Forge")]
for i, (n, c, bg, l) in enumerate(stats):
    y = 2.25 + i * 1.03
    box(s, 8.1, y, 4.83, 0.9, fill=bg, radius=0.1)
    tb(s, 8.2, y, 1.75, 0.9, [{"runs": [n], "align": "c"}], anchor="m", size=34, bold=True, color=c)
    tb(s, 9.95, y, 2.9, 0.9, [{"runs": [l]}], anchor="m", size=14.5, color=INK)
card(s, 8.1, 5.42, 4.83, 1.48, fill=GOLD_T, line=None, shadow=False)
tb(s, 8.25, 5.46, 4.6, 0.32, [{"runs": ["WHY IT MATTERS  ·  REAL-WORLD USE"]}], anchor="m", size=12.5, bold=True, color=GOLD_D)
tb(s, 8.25, 5.8, 4.6, 1.08, [{"runs": [t], "bullet": True, "after": 1} for t in
    ["SOC detection engineering and threat hunting", "Portable rules for any SIEM (Sigma export)", "Audit-ready evidence for every condition"]], size=14)
notes(s, "State the engineering problem: prose in, executable detection out is a manual, error-prone step. Three failure modes motivate the project: the manual gap, conditions that logs cannot observe, and LLM-generated rules that look right but are wrong. The three numbers on the right are real measurements from our held-out experiments, detailed on the Results slide. Only 1 of 5 direct-LLM rules and 0 of 5 schema-constrained rules were fully correct.")

# ------------------------------------------------------------------ 5. LITERATURE REVIEW
s = S[3]; prep(s, "Literature Review (Previous Work)")
lead(s, ["Prior work extracts or represents threats. ", ("None closes the loop to a tested rule.", {"bold": True, "color": MAROON})])
rows = [("TTPDrill [1]\nACSAC 2017", "NLP (POS tagging) + TF-IDF/BM25 to map report sentences to ATT&CK-style threat actions.", "Outputs technique labels only; no log-field check and no executable rule."),
        ("EXTRACTOR [2]\nEuroS&P 2021", "Extracts concise attack behaviour as provenance graphs from unstructured CTI text.", "Graphs support threat hunting; thresholds and windows never become tested rules."),
        ("AttacKG [3]\nESORICS 2022", "Technique knowledge graphs; identifies entities, dependencies and techniques (F1 0.887 / 0.896 / 0.789).", "Knowledge representation, not detection generation; no telemetry awareness."),
        ("SecureBERT [4]\nSecureComm 2022", "Domain-adapted BERT language model for cybersecurity text.", "A representation model: needs task heads and offers no rule synthesis or validation."),
        ("Sigma + ATT&CK [5], [6]", "Vendor-neutral rule format and a knowledge base of adversary techniques.", "Rules are still hand-authored; no link back to source text or observability checks."),
        ("Direct LLM rules [7], [8]\n(baseline in this work)", "Prompt an LLM to write detection code straight from the report.", "Compiles but hard-codes values and enum literals; hallucination and injection risks.")]
gs = s.shapes.add_table(len(rows) + 1, 3, Inches(0.4), Inches(1.95), Inches(12.53), Inches(4.2))
tbl = gs.table
tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
for j, w in enumerate((2.75, 5.0, 4.78)): tbl.columns[j].width = Inches(w)
def cell(c, paras, fill, anchor=MSO_ANCHOR.MIDDLE, border=LINE, **base):
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill); c.vertical_anchor = anchor
    c.margin_left = c.margin_right = Inches(0.12); c.margin_top = c.margin_bottom = Inches(0.04)
    c.text_frame.word_wrap = True
    fill_tf(c.text_frame, paras, base)
    tcPr = c._tc.get_or_add_tcPr()
    ln = etree.Element(qn("a:lnB"), w="9525"); sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=border)
    tcPr.insert(0, ln)
tbl.rows[0].height = Inches(0.42)
for j, h in enumerate(("Paper / System", "Method", "Limitation")):
    cell(tbl.cell(0, j), [{"runs": [h]}], MAROON, bold=True, size=15, font=SERIF, color=WHITE, border=MAROON)
for i, (a, b, c_) in enumerate(rows, 1):
    tbl.rows[i].height = Inches(0.63)
    bg = WHITE if i % 2 else PAPER
    l1, _, l2 = a.partition("\n")
    ps = [{"runs": [l1], "bold": True, "size": 14, "font": SERIF, "color": MAROON}]
    if l2: ps.append({"runs": [l2], "size": 11.5, "color": GREY})
    cell(tbl.cell(i, 0), ps, bg)
    cell(tbl.cell(i, 1), [{"runs": [b]}], bg, size=12.5)
    cell(tbl.cell(i, 2), [{"runs": [c_]}], bg, size=12.5)
box(s, 0.4, 6.32, 12.53, 0.58, fill=GOLD_T, radius=0.15)
icon(s, "FaLightbulb", "gold", 0.58, 6.44, 0.34)
tb(s, 1.05, 6.32, 11.8, 0.58, [{"runs": [("Research gap:  ", {"bold": True, "color": GOLD_D}), "no prior system verifies that the logs can observe a behaviour before compiling it into a tested, evidence-linked rule."]}], anchor="m", size=15)
notes(s, "Previous work falls into four groups: NLP extraction of attacker behaviour (TTPDrill, EXTRACTOR, AttacKG), domain language models (SecureBERT), standards (Sigma, ATT&CK) and direct LLM generation. Each stops before the same step: nobody checks whether the available logs can observe the behaviour, compiles it deterministically, and replay-tests the result. That gap is the project's contribution. Bracketed numbers match the References slide.")

# ------------------------------------------------------------------ 6. NOVELTY (new)
s = clone_skeleton("Novelty of the Project")
lead(s, [("Telemetry-aware compilation", {"bold": True, "color": MAROON}), " with evidence and honest refusal."])
pill = [("FaEye", "Observability gate", "Stage 3 verifies every required field exists in the documented schema, else it rejects. With the gate off, an unsupported case compiled into a rule that never fires."),
        ("FaPuzzlePiece", "Closed-recipe compiler", "Only 3 parameterised Scala/Spark recipes, so hard-coded values and enum hallucination cannot occur by construction."),
        ("FaFingerprint", "Verified provenance", "Each extracted field carries a character offset into the report; quotes are re-checked against the source (100% verified)."),
        ("FaBalanceScale", "Agreement signal", "Classical vs transformer disagreement flags a report for review; Pearson r = 0.919 with correctness (n = 5).")]
for i, (ic, t, d) in enumerate(pill):
    r, cix = divmod(i, 2)
    x = 0.4 + cix * 3.68; y = 1.95 + r * 2.2
    card(s, x, y, 3.48, 2.05)
    icon_circle(s, ic, x + 0.18, y + 0.18, 0.55)
    tb(s, x + 0.82, y + 0.14, 2.6, 0.62, [{"runs": [t]}], anchor="m", size=17, bold=True, font=SERIF, color=MAROON)
    tb(s, x + 0.18, y + 0.8, 3.14, 1.22, [{"runs": [d]}], size=13.5)
# comparison matrix
mx = 7.85; mw = 5.08
box(s, mx, 1.95, mw, 4.35, fill=WHITE, line=LINE, shadow=True, radius=0.03)
box(s, mx, 1.95, mw, 0.62, fill=MAROON, radius=0.03)
tb(s, mx + 0.1, 1.95, 2.2, 0.62, [{"runs": ["How it compares"]}], anchor="m", size=15, bold=True, font=SERIF, color=WHITE)
cols = [("Prior CTI\nextractors", 3.2), ("Direct\nLLM", 3.98), ("SENTINEL\nForge", 4.4)]
cxs = [mx + 2.3, mx + 3.3, mx + 4.13]
for (lab, _), cx_ in zip(cols, cxs):
    tb(s, cx_ - 0.2, 1.97, 1.0, 0.58, [{"runs": [lab], "align": "c"}], anchor="m", size=11.5, bold=True, color=WHITE, m=(0, 0, 0, 0))
matrix = [("Extracts behaviour from text", "y", "y", "y"), ("Emits an executable rule", "n", "y", "y"), ("Checks log observability first", "n", "n", "y"),
          ("Cites source evidence per field", "p", "n", "y"), ("Refuses unsupported behaviours", "n", "n", "y"), ("Replay-tested on labelled events", "n", "n", "y")]
for i, row in enumerate(matrix):
    y = 2.6 + i * 0.6
    if i % 2 == 0: box(s, mx + 0.03, y, mw - 0.06, 0.6, fill=PAPER, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    tb(s, mx + 0.1, y, 2.2, 0.6, [{"runs": [row[0]]}], anchor="m", size=12.5)
    for v, cx_ in zip(row[1:], cxs):
        bg, ic = {"y": (GREEN, "FaCheck"), "n": (RED, "FaTimes"), "p": (AMBER, "FaExclamation")}[v]
        icon_circle(s, ic, cx_ + 0.15, y + 0.15, 0.3, bg=bg, ratio=0.55)
box(s, 0.4, 6.42, 12.53, 0.48, fill=NAVY_T, radius=0.2)
tb(s, 0.6, 6.42, 12.2, 0.48, [{"runs": [("Scoped claim:  ", {"bold": True, "color": NAVY}), "shown on one log schema and five behaviours. Replay success is not a guarantee of catching every real attack."], "italic": False}], anchor="m", size=14)
notes(s, "Novelty is not a better language model; it is a different pipeline shape. Four mechanisms: (1) an observability gate that rejects behaviours the logs cannot see - in our ablation, removing it produced a detector that compiles, deploys and never fires; (2) a closed set of three compiler recipes, so the output space is restricted; (3) provenance with verified character offsets; (4) cross-extractor agreement as a calibrated review signal. Be precise about scope: one schema, five behaviours, small held-out sample.")

# ------------------------------------------------------------------ 7. OBJECTIVES
s = S[4]; prep(s, "Objectives of the Project")
lead(s, ["Six measurable objectives for a ", ("trustworthy", {"bold": True, "color": MAROON}), " report-to-rule pipeline."])
objs = [("FaSearch", "Extract behaviour", "Identify actions, entities, conditions and event sequences with classical and transformer NLP, compared on F1."),
        ("FaEye", "Check observability", "Decide whether the documented log schema can capture the behaviour before any rule is generated."),
        ("FaCode", "Compile precisely", "Convert a typed behaviour spec into a restricted Scala/Spark rule with exact detection semantics."),
        ("FaSyncAlt", "Validate reliability", "Replay labelled events; stress with missing fields, policy outages and adversarial text."),
        ("FaLink", "Preserve evidence", "Link every detection condition to source text or to an explicit analyst assumption."),
        ("FaBan", "Reject unsupported", "Ask for clarification or withhold generation when essential information is missing.")]
for i, (ic, t, d) in enumerate(objs):
    r, cix = divmod(i, 3)
    x = 0.4 + cix * (cw + 0.245); y = 1.95 + r * 2.2
    card(s, x, y, cw, 2.0)
    num_circle(s, i + 1, x + 0.18, y + 0.18, 0.46, size=18)
    icon(s, ic, "maroon", x + cw - 0.65, y + 0.24, 0.36)
    tb(s, x + 0.75, y + 0.14, cw - 1.5, 0.55, [{"runs": [t]}], anchor="m", size=18, bold=True, font=SERIF, color=MAROON)
    tb(s, x + 0.2, y + 0.82, cw - 0.4, 1.15, [{"runs": [d]}], size=16)
box(s, 0.4, 6.4, 12.53, 0.5, fill=ROSE, radius=0.2)
tb(s, 0.6, 6.4, 12.2, 0.5, [{"runs": [("Research question:  ", {"bold": True, "color": MAROON}), "can telemetry-aware compilation and validation beat direct LLM rule generation on correctness and reliability?"]}], anchor="m", size=14.5)
notes(s, "Objectives map one-to-one to pipeline stages. The research question is tested by comparing four systems on the same held-out reports and the same labelled events: manual rules, direct LLM generation, schema-constrained generation, and the full SENTINEL Forge pipeline.")

# ------------------------------------------------------------------ 8. METHODOLOGY
s = S[5]; prep(s, "Proposed Methodology")
lead(s, ["Five stages from raw report to ", ("explained alert", {"bold": True, "color": MAROON}), "."])
stages = [("FaDatabase", "Ingest & Normalize", "Scala/Spark cleans telemetry into a partitioned Parquet store for replay and streaming.", "Parquet store"),
          ("FaSearch", "Understand the Report", "Classical (regex) vs transformer (LLM, controlled vocabulary) extraction, with verified provenance.", "Behaviour IR (JSON)"),
          ("FaClipboardCheck", "Validate the Spec", "Stage 3 gate: fields exist and are reliable, values are sane, structural dependencies covered.", "supported / rejected"),
          ("FaCogs", "Compile & Execute", "spec_bridge → CompiledSpec → Scala RuleCompiler → Spark batch and streaming; Sigma export.", "Alerts + Sigma YAML"),
          ("FaChartLine", "Evaluate & Explain", "Compare with independent labels; dashboard shows evidence, verdict and analyst decision.", "Report + audit log")]
sw = 2.34
for i, (ic, t, d, o) in enumerate(stages):
    x = 0.4 + i * (sw + 0.2)
    card(s, x, 1.95, sw, 3.4)
    num_circle(s, i + 1, x + 0.15, 2.1, 0.46, size=18)
    icon(s, ic, "maroon", x + sw - 0.6, 2.15, 0.36)
    tb(s, x + 0.12, 2.68, sw - 0.24, 0.62, [{"runs": [t]}], anchor="m", size=16.5, bold=True, font=SERIF, color=MAROON)
    tb(s, x + 0.12, 3.32, sw - 0.24, 1.5, [{"runs": [d]}], size=13.5)
    chip(s, x + 0.12, 4.86, sw - 0.24, 0.36, o, fill=ROSE, color=MAROON, size=12)
    if i < 4: chevron(s, x + sw - 0.02, 3.4, 0.24, 0.4, color=GOLD)
card(s, 0.4, 5.52, 6.05, 1.38, fill=NAVY_T, line=None, shadow=False)
icon(s, "FaKey", "navy", 0.6, 5.7, 0.4)
tb(s, 1.15, 5.56, 5.2, 1.3, [{"runs": [("Decision rule", {"bold": True, "color": NAVY, "font": SERIF, "size": 15})], "after": 2},
    {"runs": [("supported", {"bold": True, "color": GREEN}), "  →  compile and replay"]},
    {"runs": [("unsupported / ambiguous", {"bold": True, "color": RED}), "  →  reject with a reason or ask the analyst. Never guess."]}], size=14)
tb(s, 6.6, 5.5, 6.33, 0.3, [{"runs": ["THREE CLOSED RECIPES  (the only rules the compiler can emit)"]}], anchor="m", size=12, bold=True, color=GREY)
rec = [("SequenceThenTrigger", "B1 · brute force"), ("DistinctCountWithinWindow", "B2 spray · B3 sessions"), ("PolicyCompare", "B4 · B5 policy checks")]
rw = [1.85, 2.4, 1.72]
xx = 6.6
for (n, sub), w in zip(rec, rw):
    box(s, xx, 5.84, w, 1.02, fill=WHITE, line=MAROON, lw=1.25, radius=0.1)
    tb(s, xx + 0.05, 5.86, w - 0.1, 0.55, [{"runs": [n], "align": "c"}], anchor="m", size=11.5, bold=True, font=MONO, color=MAROON, m=(0.02, 0, 0.02, 0))
    tb(s, xx + 0.05, 6.4, w - 0.1, 0.4, [{"runs": [sub], "align": "c"}], anchor="m", size=12, color=GREY, m=(0.02, 0, 0.02, 0))
    xx += w + 0.12
notes(s, "Five stages. Stage 1 normalises telemetry. Stage 2 extracts a typed behaviour specification with two competing extractors so we can measure classical versus transformer. Stage 3 is the observability gate. Stage 4 turns an approved spec into Spark logic via a closed set of three recipes. Stage 5 compares alerts to independent labels and presents evidence to an analyst. The guiding rule: if something cannot be supported, the system says so instead of guessing.")

# ------------------------------------------------------------------ 9. IMPLEMENTATION LAYOUT (new)
s = clone_skeleton("Implementation Layout")
lead(s, ["Three layers, connected by ", ("two JSON contracts", {"bold": True, "color": MAROON}), "."])
bands = [
 (MAROON, ROSE, "FaPython", "Python", "Understand & Validate",
  [("nlp/classical_extractor.py", "regex / keyword baseline"), ("nlp/transformer_extractor.py", "LLM, controlled vocabulary"), ("nlp/injection_guard.py", "blocks prompt injection"),
   ("nlp/schema_fields.py", "single source of vocabulary"), ("compiler/observability_checker.py", "Stage 3 gate"), ("compiler/spec_bridge.py", "extraction → compiled spec")]),
 (NAVY, NAVY_T, "FaCogs", "Scala + Spark", "Compile & Execute",
  [("CompiledSpec.scala", "typed spec + loader"), ("RuleCompiler.scala", "closed 3-recipe compiler"), ("ReplayCheck.scala", "batch replay vs labels"), ("StreamingCheck.scala", "Structured Streaming"),
   ("RobustnessCheck.scala", "outage / missing field"), ("ThroughputCheck.scala", "latency, events/s"), ("EndToEndCheck.scala", "extraction-derived specs"), ("sigma_export.py", "Sigma YAML export")]),
 (GOLD_D, GOLD_T, "FaDatabase", "Data · Review", "Evidence & dashboard",
  [("data/", "reports · Parquet · labels"), ("docs/", "schema · IR · semantics"), ("experiments/", "baselines · ablations · results"), ("compiler/test/", "25 regression cases"), ("dashboard/", "FastAPI + vanilla JS")])]
by = 1.95; bh = 1.42; bg_ = 0.32
def flow_chips(slide, items, x0, y0, maxw, border):
    x, y = x0, y0
    for name, role in items:
        w = max(len(name) * 0.083 + 0.28, len(role) * 0.068 + 0.28)
        if x + w > x0 + maxw: x = x0; y += 0.62
        box(slide, x, y, w, 0.54, fill=WHITE, line=border, lw=1.0, radius=0.12)
        tb(slide, x, y + 0.02, w, 0.28, [{"runs": [name], "align": "c"}], anchor="m", size=11, bold=True, font=MONO, color=border, m=(0.02, 0, 0.02, 0))
        tb(slide, x, y + 0.29, w, 0.24, [{"runs": [role], "align": "c"}], anchor="m", size=10.5, color=GREY, m=(0.02, 0, 0.02, 0))
        x += w + 0.1
for i, (col, tint, ic, t1, t2, chips) in enumerate(bands):
    y = by + i * (bh + bg_)
    box(s, 0.4, y, 12.53, bh, fill=tint, radius=0.05)
    box(s, 0.4, y, 2.5, bh, fill=col, radius=0.05)
    icon(s, ic, "white", 0.6, y + 0.32, 0.44)
    tb(s, 1.12, y + 0.22, 1.75, 0.6, [{"runs": [t1]}], anchor="m", size=17, bold=True, font=SERIF, color=WHITE)
    tb(s, 0.6, y + 0.85, 2.25, 0.5, [{"runs": [t2]}], size=13.5, color=WHITE)
    flow_chips(s, chips, 3.1, y + 0.1, 9.7, col)
for i, lab in enumerate(["Compiled spec  (single-line JSON)", "Alerts · verdicts · results JSON"]):
    y = by + (i + 1) * bh + i * bg_ + (i) * 0 + i * 0
    yy = by + (i + 1) * bh + i * bg_
    box(s, 6.15, yy + 0.04, 0.28, 0.24, fill=MAROON, shape=MSO_SHAPE.DOWN_ARROW)
    tb(s, 6.5, yy, 3.6, 0.32, [{"runs": [lab]}], anchor="m", size=12, bold=True, color=MAROON)
notes(s, "This is how the repository is laid out. Top layer, Python: two extractors, a prompt-injection guard, the Stage 3 observability checker and the spec bridge that converts a validated extraction into a compiled spec. Middle layer, Scala and Spark: the closed-recipe RuleCompiler plus separate check programs for replay, streaming, robustness, throughput and end-to-end. Bottom layer: data, docs, experiments and the FastAPI dashboard. The two JSON contracts are the Behaviour IR and the Compiled Spec.")

# ------------------------------------------------------------------ 10. TOOLS
s = S[6]; prep(s, "Tools & Technologies")
lead(s, ["A Python NLP front end feeding a ", ("Scala/Spark", {"bold": True, "color": MAROON}), " compiler and execution engine."])
cols3 = [
 (MAROON, "FaMicrochip", "Software", [("Python 3", "NLP, validation, dashboard API"), ("Scala 2.13.12", "restricted-rule compiler"), ("Apache Spark 3.5.3 [9]", "Spark SQL, Structured Streaming [10]"),
   ("sbt + ScalaTest", "build and test"), ("FastAPI + Pydantic", "analyst review API"), ("HTML / CSS / JS", "no-build review dashboard"), ("Groq · gpt-oss-120b", "pretrained transformer, not fine-tuned")]),
 (NAVY, "FaServer", "Hardware & Environment", [("Single local machine", "Spark local[*] on the JVM"), ("No GPU required", "transformer is called via API"), ("Dashboard on localhost", "served by Uvicorn"), ("Cluster deployment", "planned, see Future Work")]),
 (GOLD_D, "FaLayerGroup", "Data & Standards", [("Partitioned Parquet", "event store + Spark SQL replay"), ("MITRE ATT&CK [5]", "T1110 · T1078 · T1556 · T1621"), ("Sigma YAML [6]", "portable rule export"),
   ("JSON contracts", "Behaviour IR · Compiled spec"), ("Auth-log schema v1", "9 fields + 1 policy table")])]
for i, (col, ic, t, its) in enumerate(cols3):
    x = 0.4 + i * (cw + 0.245)
    card(s, x, 1.95, cw, 4.95)
    box(s, x, 1.95, cw, 0.7, fill=col, radius=0.05)
    icon(s, ic, "white", x + 0.2, 2.1, 0.4)
    tb(s, x + 0.75, 1.95, cw - 0.85, 0.7, [{"runs": [t]}], anchor="m", size=19, bold=True, font=SERIF, color=WHITE)
    step = 4.15 / max(len(its), 5) if len(its) > 4 else 0.95
    for j, (a, b) in enumerate(its):
        yy = 2.78 + j * (0.6 if len(its) > 5 else 0.72 if len(its) == 5 else 0.85)
        box(s, x + 0.2, yy + 0.15, 0.12, 0.12, fill=col, shape=MSO_SHAPE.OVAL)
        tb(s, x + 0.42, yy, cw - 0.55, 0.58, [{"runs": [(a, {"bold": True, "size": 15})], "after": 0}, {"runs": [b], "size": 12.5, "color": GREY}], m=(0.02, 0, 0.02, 0))
box(s, 0.4 + (cw + 0.245) + 0.2, 6.02, cw - 0.4, 0.75, fill=GOLD_T, radius=0.1)
tb(s, 0.4 + (cw + 0.245) + 0.3, 6.02, cw - 0.6, 0.75, [{"runs": [("Scope note: ", {"bold": True, "color": GOLD_D}), "all measurements are single-machine; cluster runs are future work."]}], anchor="m", size=13.5)
notes(s, "Software: Python for extraction and validation, Scala 2.13 with Apache Spark 3.5.3 for the compiler and execution, sbt/ScalaTest for builds, FastAPI for the dashboard API. The transformer is a pretrained model reached through the Groq API; we did not fine-tune it because ten training reports would only memorise. Hardware: everything runs on one local machine, so throughput numbers are small-scale and honestly labelled. Standards: ATT&CK technique IDs, Sigma export and JSON contracts.")

# ------------------------------------------------------------------ 11. EXPERIMENTAL SETUP
s = S[7]; prep(s, "Experimental Setup")
lead(s, ["Same reports, same events, same labels for ", ("every", {"bold": True, "color": MAROON}), " system."])
tiles = [("15", "reports, 3 per behaviour"), ("5", "ATT&CK-mapped behaviours"), ("10 | 5", "train | held-out reports"), ("48", "replay events (Parquet)"), ("17", "labelled replay scenarios")]
tw = (12.53 - 4 * 0.2) / 5
for i, (n, l) in enumerate(tiles):
    x = 0.4 + i * (tw + 0.2)
    card(s, x, 1.95, tw, 1.2, fill=ROSE, line=None, shadow=False)
    tb(s, x, 1.98, tw, 0.68, [{"runs": [n], "align": "c"}], anchor="m", size=32, bold=True, font=SERIF, color=MAROON)
    tb(s, x, 2.66, tw, 0.42, [{"runs": [l], "align": "c"}], anchor="m", size=13, color=INK)
card(s, 0.4, 3.35, 6.2, 3.55)
tb(s, 0.6, 3.42, 5.8, 0.42, [{"runs": ["Systems compared"]}], anchor="m", size=19, bold=True, font=SERIF, color=MAROON)
sysl = [("Manual rules", "5 hand-written Scala rules (the fair baseline)"), ("Direct LLM", "report → Scala rule in one shot"),
        ("Schema-constrained LLM", "same, limited to schema fields; may reject"), ("SENTINEL Forge", "NLP → Stage 3 → compiler → Spark, plus a gate-off ablation")]
for i, (a, b) in enumerate(sysl):
    y = 3.92 + i * 0.72
    num_circle(s, i + 1, 0.62, y + 0.06, 0.4, bg=MAROON if i == 3 else NAVY, size=15)
    tb(s, 1.15, y, 5.3, 0.62, [{"runs": [(a, {"bold": True, "size": 15.5})], "after": 0}, {"runs": [b], "size": 13, "color": GREY}], anchor="m", m=(0.02, 0, 0.02, 0))
card(s, 6.8, 3.35, 6.13, 3.55)
tb(s, 7.0, 3.42, 5.8, 0.42, [{"runs": ["What is measured, and how it stays fair"]}], anchor="m", size=19, bold=True, font=SERIF, color=MAROON)
meas = [("Extraction:", " precision, recall, F1, behaviour accuracy, provenance rate"), ("Detection:", " replay verdicts vs independent labels, incl. window-boundary negatives"),
        ("Robustness:", " policy outage, missing field, 6 adversarial fixtures"), ("Performance:", " 10 runs per behaviour: mean, p50, p95, events/s"),
        ("Leakage control:", " paraphrases stay in train; held-out = a different incident")]
tb(s, 7.0, 3.9, 5.8, 3.0, [{"runs": [(a, {"bold": True, "color": NAVY}), b], "bullet": True, "after": 9} for a, b in meas], size=15.5)
notes(s, "Dataset: 15 synthetic-but-realistic threat reports, three per behaviour. Ten are for training and prompt development; five held-out reports are each a different incident, not a paraphrase, to avoid leakage. Events: a 48-event partitioned Parquet store with 17 independently authored labelled scenarios, including exact window-boundary negatives. The log schema has nine fields; source_ip is deliberately marked unreliable so the gate has a real case to reason about.")

# ------------------------------------------------------------------ 12. RESULTS 1/2
s = S[8]; prep(s, "Results (1/2): Detection Correctness")
lead(s, [("SENTINEL Forge is correct on every replayed scenario", {"bold": True, "color": MAROON}), "; LLM-only baselines are not."])
mx0 = 0.4; lw_ = 2.4; bwid = 0.84; sw_ = 1.3
box(s, mx0, 1.95, lw_ + 5 * bwid + sw_, 0.72, fill=MAROON, radius=0.04)
heads = ["B1\nbrute force", "B2\nspray", "B3\nsessions", "B4\nsvc acct", "B5\nMFA"]
tb(s, mx0 + 0.1, 1.95, lw_, 0.72, [{"runs": ["System"]}], anchor="m", size=15, bold=True, font=SERIF, color=WHITE)
for i, h in enumerate(heads):
    a, b = h.split("\n")
    tb(s, mx0 + lw_ + i * bwid, 1.95, bwid, 0.72, [{"runs": [a], "align": "c", "size": 13.5, "bold": True}, {"runs": [b], "align": "c", "size": 10.5}], anchor="m", color=WHITE, m=(0, 0, 0, 0))
tb(s, mx0 + lw_ + 5 * bwid, 1.95, sw_, 0.72, [{"runs": ["Result"], "align": "c"}], anchor="m", size=15, bold=True, font=SERIF, color=WHITE)
res = [("Manual rules", "yyyyy", "17/17", "scenarios"), ("Direct LLM", "ynnnp", "1/5", "behaviours"), ("Schema-constrained", "nnnnp", "0/5", "behaviours"), ("SENTINEL Forge", "yyyyy", "17/17", "scenarios")]
# reorder symbols: direct LLM = B1 ok, B2 partial, B3 no, B4 no, B5 no ; schema = B1 no, B2 no, B3 no, B4 no, B5 partial
res[1] = ("Direct LLM", "ypnnn", "1/5", "behaviours")
for r, (nm, syms, sc, unit) in enumerate(res):
    y = 2.72 + r * 0.82
    hi = nm.startswith("SENT")
    box(s, mx0, y, lw_ + 5 * bwid + sw_, 0.76, fill=GREEN_T if hi else (PAPER if r % 2 == 0 else WHITE), line=GREEN if hi else LINE, lw=1.5 if hi else 0.5, radius=0.05)
    tb(s, mx0 + 0.1, y, lw_, 0.76, [{"runs": [nm]}], anchor="m", size=16, bold=True, font=SERIF, color=GREEN if hi else INK)
    for i, ch in enumerate(syms):
        bg, ic = {"y": (GREEN, "FaCheck"), "n": (RED, "FaTimes"), "p": (AMBER, "FaExclamation")}[ch]
        icon_circle(s, ic, mx0 + lw_ + i * bwid + (bwid - 0.4) / 2, y + 0.18, 0.4, bg=bg, ratio=0.52)
    tb(s, mx0 + lw_ + 5 * bwid, y, sw_, 0.76, [{"runs": [sc], "align": "c", "size": 17, "bold": True, "color": GREEN if sc == "17/17" else RED}, {"runs": [unit], "align": "c", "size": 10.5, "color": GREY}], anchor="m", m=(0, 0, 0, 0))
# legend
lx = 0.5
for bg, ic, lab in ((GREEN, "FaCheck", "fully correct"), (AMBER, "FaExclamation", "partly correct (duplicates / wrong extras)"), (RED, "FaTimes", "missed or wrong")):
    icon_circle(s, ic, lx, 6.12, 0.26, bg=bg, ratio=0.55)
    wtxt = len(lab) * 0.075 + 0.2
    tb(s, lx + 0.32, 6.07, wtxt, 0.36, [{"runs": [lab]}], anchor="m", size=12, color=GREY)
    lx += 0.32 + wtxt + 0.15
tb(s, 0.4, 6.48, 7.9, 0.45, [{"runs": ["Executed for real on the same Parquet store and independent labels. Manual rules first scored 15/17 until a timestamp bug was fixed."], "italic": True}], anchor="m", size=12, color=GREY)
# right side
rx = 8.6; rw_ = 4.33
kp = [("17/17", "replay scenarios (compiler proven on all 5 behaviours)"), ("11/11", "end-to-end: extraction-derived specs, 3 reports"), ("2/2", "robustness: policy outage and missing field handled safely")]
for i, (n, l) in enumerate(kp):
    y = 1.95 + i * 0.93
    box(s, rx, y, rw_, 0.83, fill=GREEN_T, radius=0.1)
    tb(s, rx + 0.08, y, 1.45, 0.83, [{"runs": [n], "align": "c"}], anchor="m", size=28, bold=True, color=GREEN)
    tb(s, rx + 1.55, y, rw_ - 1.65, 0.83, [{"runs": [l]}], anchor="m", size=13, m=(0.02, 0, 0.02, 0))
card(s, rx, 4.8, rw_, 2.1, fill=RED_T, line=None, shadow=False)
tb(s, rx + 0.15, 4.83, rw_ - 0.3, 0.34, [{"runs": ["WHY THE BASELINES FAILED"]}], anchor="m", size=12, bold=True, color=RED)
tb(s, rx + 0.15, 5.17, rw_ - 0.3, 1.7, [{"runs": [t], "bullet": True, "after": 2, "bcolor": RED} for t in
   ["Hard-coded the one account from the source report (B4, B5)", "Wrong enum: “failure” instead of login_failure", "7 duplicate alerts for a single incident (B2)", "Gate off: a rule that compiles and never fires"]], size=12.5)
notes(s, "This is the core result. Each system was executed against the same 48-event Parquet store and independent labels. SENTINEL Forge and the corrected manual rules pass all 17 scenarios. Only one direct-LLM rule (B1) was fully usable; schema-constrained generation produced zero fully correct rules. Important honesty point from an independent review: the original 17/17 used hand-authored compiled specs, so we built the extraction-to-compiler bridge and re-ran end to end on extraction-derived specs, scoring 11/11 on three reports.")

# ------------------------------------------------------------------ 13. RESULTS 2/2 (new)
s = clone_skeleton("Results (2/2): Extraction & Performance")
lead(s, ["Extraction is strong but not perfect, and ", ("the system says so", {"bold": True, "color": MAROON}), "."])
card(s, 0.4, 1.95, 6.3, 4.15)
tb(s, 0.6, 2.0, 5.9, 0.4, [{"runs": ["Field extraction on 5 held-out reports"]}], anchor="m", size=17, bold=True, font=SERIF, color=MAROON)
cd = CategoryChartData(); cd.categories = ["Precision", "Recall", "F1"]
cd.add_series("Classical (regex)", (1.0, 0.947, 0.973)); cd.add_series("Transformer (LLM)", (1.0, 0.895, 0.944))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.55), Inches(2.4), Inches(6.0), Inches(2.75), cd)
ch = gf.chart; ch.font.size = Pt(12); ch.font.name = SANS
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(12)
pl = ch.plots[0]; pl.gap_width = 70; pl.overlap = -8; pl.has_data_labels = True
dl = pl.data_labels; dl.number_format = "0.000"; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(12); dl.font.bold = True
for ser, col in zip(pl.series, (MAROON, NAVY)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
va = ch.value_axis; va.minimum_scale = 0; va.maximum_scale = 1.2; va.visible = False; va.has_major_gridlines = False
ca = ch.category_axis; ca.tick_labels.font.size = Pt(13); ca.format.line.color.rgb = rgb(LINE)
ms = [("0.80 → 1.00", "behaviour accuracy, classical → transformer"), ("100%", "provenance quotes verified against source")]
for i, (n, l) in enumerate(ms):
    x = 0.6 + i * 3.05
    box(s, x, 5.25, 2.9, 0.75, fill=ROSE, radius=0.1)
    tb(s, x, 5.25, 1.25, 0.75, [{"runs": [n], "align": "c"}], anchor="m", size=15 if i == 0 else 20, bold=True, color=MAROON, m=(0.02, 0, 0.02, 0))
    tb(s, x + 1.25, 5.25, 1.6, 0.75, [{"runs": [l]}], anchor="m", size=11, color=INK, m=(0.02, 0, 0.02, 0))
card(s, 6.95, 1.95, 5.98, 2.85)
tb(s, 7.15, 2.0, 5.6, 0.4, [{"runs": ["Robustness to adversarial text"]}], anchor="m", size=17, bold=True, font=SERIF, color=MAROON)
adv = [("Synonym substitution", "Classical collapsed to 1 field; transformer kept the full spec.", "p"), ("Alarming tone, compliant facts", "Classical misclassified; transformer read through the tone.", "p"),
       ("Contradictory threshold (5 → 10)", "Both extractors resolved to the corrected value.", "y"), ("Prompt injection (3 fixtures)", "Guard blocks the LLM call; constrained prompt also held.", "y")]
for i, (a, b, v) in enumerate(adv):
    y = 2.45 + i * 0.58
    bg, ic = {"y": (GREEN, "FaCheck"), "p": (AMBER, "FaExclamation")}[v]
    icon_circle(s, ic, 7.15, y + 0.08, 0.36, bg=bg, ratio=0.5)
    tb(s, 7.65, y, 5.2, 0.56, [{"runs": [(a, {"bold": True, "size": 13.5})], "after": 0}, {"runs": [b], "size": 12, "color": GREY}], anchor="m", m=(0.02, 0, 0.02, 0))
perf = [("503 ms", "mean per query · p95 1136 ms"), ("≈ 95/s", "events per second on 48 events"), ("r = 0.919", "agreement ↔ correctness (n = 5)")]
pw = (5.98 - 2 * 0.15) / 3
for i, (n, l) in enumerate(perf):
    x = 6.95 + i * (pw + 0.15)
    box(s, x, 4.95, pw, 1.15, fill=NAVY_T, radius=0.08)
    tb(s, x, 4.98, pw, 0.55, [{"runs": [n], "align": "c"}], anchor="m", size=22, bold=True, color=NAVY)
    tb(s, x + 0.05, 5.5, pw - 0.1, 0.55, [{"runs": [l], "align": "c"}], anchor="t", size=11.5, color=INK, m=(0.02, 0, 0.02, 0))
box(s, 0.4, 6.28, 12.53, 0.62, fill=GOLD_T, radius=0.15)
icon(s, "FaExclamationTriangle", "gold", 0.6, 6.42, 0.32)
tb(s, 1.1, 6.28, 11.7, 0.62, [{"runs": [("Read with care:  ", {"bold": True, "color": GOLD_D}), "5 reports, 48 events, one machine: directional results, not statistically significant or production scale."]}], anchor="m", size=14)
notes(s, "Left: field extraction. Interestingly, the regex baseline edges out the transformer on field F1 (0.973 vs 0.944) because the reports follow a template; but it misclassified one behaviour (accuracy 0.80 vs 1.00) and collapses under paraphrase. The transformer is robust to synonyms and tone. All extracted evidence quotes were verified against the source text. Right: performance was measured on a 48-event dataset on one machine; it is dominated by Spark planning overhead and is not a scale claim. Agreement between the two extractors correlated with correctness at r = 0.919 on five reports; that is why the dashboard flags disagreement for analyst review.")

# ------------------------------------------------------------------ 14. CHALLENGES
s = S[9]; prep(s, "Challenges")
lead(s, ["Every hard problem here was found by ", ("running", {"bold": True, "color": MAROON}), " the system, not by reading code."])
chs = [("FaClock", "Silent timestamp bug", "Casting an ISO string to long returned null, so time windows became unbounded. Manual rules scored 15/17, then 17/17 after the fix.", "Fixed"),
       ("FaRedo", "Extractor variance", "Transformer output differs run to run even at temperature 0; some policy fields get dropped, so a clean 5/5 first try is not guaranteed.", "Open"),
       ("FaUserCheck", "Independent audit findings", "Results ran on hand-authored specs; the validator accepted empty or invalid specs. Bridge built, validator hardened, 25 regression tests.", "Fixed"),
       ("FaStream", "Streaming limits", "Spark cannot run these window functions on streams: only PolicyCompare streams. Others need a watermark-based rewrite.", "Open"),
       ("FaBalanceScale", "Semantic edge cases", "Repeated incidents collapse into one alert for a host; a NULL log value is read as compliant instead of insufficient context.", "Open"),
       ("FaDatabase", "Data and scale limits", "15 reports, 48 events, one machine: F1 and latency are small-sample numbers, and failure recovery at scale is untested.", "Open")]
for i, (ic, t, d, st) in enumerate(chs):
    r, cix = divmod(i, 3)
    x = 0.4 + cix * (cw + 0.245); y = 1.95 + r * 2.5
    card(s, x, y, cw, 2.3)
    icon_circle(s, ic, x + 0.18, y + 0.16, 0.54)
    tb(s, x + 0.82, y + 0.14, cw - 1.9, 0.6, [{"runs": [t]}], anchor="m", size=16.5, bold=True, font=SERIF, color=MAROON)
    chip(s, x + cw - 0.95, y + 0.27, 0.78, 0.3, st, fill=GREEN if st == "Fixed" else AMBER, color=WHITE, size=11.5)
    tb(s, x + 0.18, y + 0.85, cw - 0.36, 1.4, [{"runs": [d]}], size=15)
notes(s, "Be candid here: this is a strength of the project. The timestamp bug hid for four phases because the code had never been executed. An independent review found the extraction-to-compiler wiring was missing, and we fixed it and documented it. Open items are stated plainly: extraction variance, streaming limits for two of three recipes, two semantic gaps found by a dedicated semantics check, and the small dataset.")

# ------------------------------------------------------------------ 15. TIMELINE (new)
s = clone_skeleton("Project Timeline")
lead(s, ["Seven build phases and a hardening pass ", ("completed", {"bold": True, "color": GREEN}), "; final evaluation next."])
tl = [("Phase 0", "Scope", "One schema, five behaviours, canonical example.", "done"), ("Phase 1", "Data & baselines", "15 reports, replay labels, manual and LLM baselines.", "done"),
      ("Phase 2", "NLP extraction", "Classical vs transformer, F1 and provenance.", "done"), ("Phase 3", "Validation gate", "Stage 3 observability checker, 17 cases.", "done"),
      ("Phase 4", "Compiler & Spark", "Scala recipes, replay, streaming, Sigma.", "done"), ("Phase 5", "Evaluation", "Baselines, ablation, robustness, throughput.", "done"),
      ("Phase 6", "Dashboard", "Analyst review, injection guard, audit log.", "done"), ("Hardening", "16–17 Sep 2026", "Independent review: 7 findings, bridge, 25 tests.", "done"),
      ("Next", "Final evaluation", "Stabilise extraction, larger held-out set, cluster runs.", "plan")]
n = len(tl); x0 = 1.7; x1 = 11.63; ly = 4.4
box(s, 1.2, ly - 0.03, 10.93, 0.06, fill=MAROON, shape=MSO_SHAPE.RECTANGLE)
for i, (p, t, d, st) in enumerate(tl):
    cx_ = x0 + i * (x1 - x0) / (n - 1)
    col = GREEN if st == "done" else AMBER
    box(s, cx_ - 0.19, ly - 0.19, 0.38, 0.38, fill=WHITE, line=col, lw=3, shape=MSO_SHAPE.OVAL)
    if st == "done": icon(s, "FaCheck", "green", cx_ - 0.1, ly - 0.1, 0.2)
    else: icon(s, "FaHourglassHalf", "gold", cx_ - 0.1, ly - 0.1, 0.2)
    cwid = 2.3; up = (i % 2 == 0)
    cxl = min(max(cx_ - cwid / 2, 0.4), 12.93 - cwid)
    cy = 1.95 if up else 4.85
    ch_ = 2.05
    card(s, cxl, cy, cwid, ch_, fill=WHITE if st == "done" else GOLD_T, line=LINE if st == "done" else GOLD, radius=0.06)
    chip(s, cxl + 0.14, cy + 0.12, 1.1, 0.28, p, fill=MAROON if st == "done" else GOLD_D, color=WHITE, size=11.5)
    tb(s, cxl + 0.12, cy + 0.44, cwid - 0.24, 0.42, [{"runs": [t]}], anchor="m", size=16, bold=True, font=SERIF, color=MAROON if st == "done" else GOLD_D)
    tb(s, cxl + 0.12, cy + 0.88, cwid - 0.24, 1.15, [{"runs": [d]}], size=13)
    ty = cy + ch_ if up else cy
    box(s, cx_ - 0.01, (cy + ch_) if up else (ly + 0.19), 0.02, (ly - 0.19) - (cy + ch_) if up else (cy - (ly + 0.19)), fill=col, shape=MSO_SHAPE.RECTANGLE)
notes(s, "The project ran as seven phases: scope, data and baselines, NLP extraction, validation gate, compiler and Spark execution, evaluation, and the analyst dashboard. After the build, an independent review on 16 and 17 September 2026 led to a hardening pass: seven findings reproduced and fixed, the extraction-to-compiler bridge built and run end to end, and 25 regression tests added. Next is final evaluation: stabilise extraction reliability, enlarge the held-out set and run on a real cluster. Add exact calendar dates for Phases 0-6 if your course requires them.")

# ------------------------------------------------------------------ 16. CONCLUSIONS
s = S[10]; prep(s, "Conclusions")
lead(s, ["A working, tested report-to-detection path, ", ("with its limits stated openly", {"bold": True, "color": MAROON}), "."])
conc = [("FaFlagCheckered", "Work completed", MAROON, ["All 7 phases built; 5 behaviours run end to end", "17/17 replay scenarios; 11/11 on extraction-derived specs", "25 regression tests; live analyst dashboard"]),
        ("FaRocket", "Expected outcome and impact", NAVY, ["Analysts get evidence-linked, reviewable, refusable rules", "Beats direct LLM generation on the tested set", "Not a guarantee against every real attack"]),
        ("FaTasks", "Remaining before final evaluation", GOLD_D, ["Stabilise transformer extraction (multi-run voting)", "Re-run all 5 reports cleanly, end to end", "Close the two semantic gaps"]),
        ("FaLightbulb", "Improvements planned", GREEN, ["Larger corpus and a fine-tuned domain model", "Streaming for every recipe", "Multi-node Spark cluster runs"])]
for i, (ic, t, col, bl) in enumerate(conc):
    r, cix = divmod(i, 2)
    x = 0.4 + cix * 6.38; y = 1.95 + r * 2.5
    card(s, x, y, 6.15, 2.3)
    icon_circle(s, ic, x + 0.2, y + 0.18, 0.6, bg=col)
    tb(s, x + 0.98, y + 0.18, 5.0, 0.6, [{"runs": [t]}], anchor="m", size=19, bold=True, font=SERIF, color=col)
    tb(s, x + 0.25, y + 0.92, 5.7, 1.35, [{"runs": [b], "bullet": True, "after": 3, "bcolor": col} for b in bl], size=16.5)
notes(s, "Summarise: the project delivers a complete, executed pipeline from report text to replay-verified detection for five behaviours, with an analyst dashboard. The expected impact is faster, more auditable detection engineering. Be explicit that results are on one schema, five behaviours and a small dataset. Remaining work before the final evaluation is mostly reliability and scale, listed on the next slide.")

# ------------------------------------------------------------------ 17. FUTURE WORK (new)
s = clone_skeleton("Future Work & Scope")
lead(s, ["From a verified prototype to a ", ("scalable", {"bold": True, "color": MAROON}), " detection-engineering assistant."])
fw = [(MAROON, "FaBolt", "Near term", "before final evaluation",
       [("Extraction reliability", "majority vote over N runs; alert on dropped policy fields"), ("Per-incident alerts", "fix repeated-incident collapse in the distinct-count recipe"),
        ("NULL-aware policy check", "unobserved values become insufficient_context"), ("Reference detector", "independent implementation for differential boundary tests")]),
      (NAVY, "FaRoute", "Mid term", "next iteration",
       [("Larger corpus", "many more reports per behaviour, significance-tested F1"), ("Fine-tuned domain model", "SecureBERT-class encoder [4] trained on real reports"),
        ("Streaming for all recipes", "watermarked time-window aggregations"), ("Multi-node Spark cluster", "throughput at scale and failure-recovery tests")]),
      (GOLD_D, "FaGlobe", "Long term", "scope extension",
       [("More schemas and behaviours", "DNS, endpoint and cloud audit logs"), ("Validate Sigma exports", "on real SIEM back ends"),
        ("Analyst feedback loop", "approve / refine decisions improve extraction"), ("Live report feeds", "ingestion with drift monitoring")])]
for i, (col, ic, t, sub, its) in enumerate(fw):
    x = 0.4 + i * (cw + 0.245)
    card(s, x, 1.95, cw, 4.95)
    box(s, x, 1.95, cw, 0.85, fill=col, radius=0.05)
    icon(s, ic, "white", x + 0.2, 2.15, 0.45)
    tb(s, x + 0.8, 1.98, cw - 0.9, 0.5, [{"runs": [t]}], anchor="m", size=21, bold=True, font=SERIF, color=WHITE)
    tb(s, x + 0.8, 2.42, cw - 0.9, 0.32, [{"runs": [sub]}], anchor="m", size=12.5, color=WHITE, italic=True)
    for j, (a, b) in enumerate(its):
        yy = 2.98 + j * 0.97
        num_circle(s, j + 1, x + 0.2, yy + 0.12, 0.34, bg=col, size=13)
        tb(s, x + 0.66, yy + 0.02, cw - 0.8, 0.93, [{"runs": [(a, {"bold": True, "size": 15.5})], "after": 0}, {"runs": [b], "size": 13, "color": GREY}], anchor="t", m=(0.02, 0, 0.02, 0))
notes(s, "Future work in three horizons. Near term fixes the reliability and semantic gaps already identified by our own tests. Mid term scales the data, the model and the cluster. Long term widens scope to more log sources and closes the loop with analyst feedback. Every near-term item is a named, reproducible gap from the Challenges slide, not a vague aspiration.")

# ------------------------------------------------------------------ 18. REFERENCES
s = S[11]
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.startswith("(Use IEEE"): sh._element.getparent().remove(sh._element)
body = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("R. L. Siegel")][0]
refs = [
 [("G. Husari, E. Al-Shaer, M. Ahmed, B. Chu, and X. Niu, “TTPDrill: Automatic and accurate extraction of threat actions from unstructured text of CTI sources,” in ", 0), ("Proc. 33rd Annu. Comput. Security Appl. Conf. (ACSAC)", 1), (", 2017, pp. 103–115.", 0)],
 [("K. Satvat, R. Gjomemo, and V. N. Venkatakrishnan, “EXTRACTOR: Extracting attack behavior from threat reports,” in ", 0), ("Proc. IEEE Eur. Symp. Security Privacy (EuroS&P)", 1), (", 2021.", 0)],
 [("Z. Li, J. Zeng, Y. Chen, and Z. Liang, “AttacKG: Constructing technique knowledge graph from cyber threat intelligence reports,” in ", 0), ("Computer Security – ESORICS 2022", 1), (". Cham, Switzerland: Springer, 2022, doi: 10.1007/978-3-031-17140-6_29.", 0)],
 [("E. Aghaei, X. Niu, W. Shadid, and E. Al-Shaer, “SecureBERT: A domain-specific language model for cybersecurity,” in ", 0), ("Security and Privacy in Communication Networks (SecureComm 2022)", 1), (". Springer, 2023, pp. 39–56.", 0)],
 [("B. E. Strom ", 0), ("et al.", 1), (", “MITRE ATT&CK: Design and philosophy,” The MITRE Corporation, Tech. Rep., 2018.", 0)],
 [("SigmaHQ, “Sigma: Generic signature format for SIEM systems.” GitHub. [Online]. Available: https://github.com/SigmaHQ/sigma", 0)],
 [("Z. Ji ", 0), ("et al.", 1), (", “Survey of hallucination in natural language generation,” ", 0), ("ACM Comput. Surv.", 1), (", vol. 55, no. 12, Art. 248, 2023.", 0)],
 [("K. Greshake ", 0), ("et al.", 1), (", “Not what you’ve signed up for: Compromising real-world LLM-integrated applications with indirect prompt injection,” in ", 0), ("Proc. 16th ACM Workshop Artif. Intell. Security (AISec)", 1), (", 2023, pp. 79–90.", 0)],
 [("M. Zaharia ", 0), ("et al.", 1), (", “Apache Spark: A unified engine for big data processing,” ", 0), ("Commun. ACM", 1), (", vol. 59, no. 11, pp. 56–65, 2016.", 0)],
 [("M. Armbrust ", 0), ("et al.", 1), (", “Structured Streaming: A declarative API for real-time applications in Apache Spark,” in ", 0), ("Proc. ACM SIGMOD Int. Conf. Manage. Data", 1), (", 2018, pp. 601–613.", 0)]]
txBody = body.text_frame._txBody
paras = txBody.findall(qn("a:p"))
tmpl_p = copy.deepcopy(paras[1])       # 2nd paragraph: plain numbered, justified
tmpl_r = copy.deepcopy(tmpl_p.find(qn("a:r")))
for p in paras: txBody.remove(p)
FS = "1500"
for ref in refs:
    p = copy.deepcopy(tmpl_p)
    for r in p.findall(qn("a:r")): p.remove(r)
    end = p.find(qn("a:endParaRPr"))
    for txt, it in ref:
        r = copy.deepcopy(tmpl_r); rPr = r.find(qn("a:rPr")); rPr.set("sz", FS)
        if it: rPr.set("i", "1")
        r.find(qn("a:t")).text = txt
        (end.addprevious(r) if end is not None else p.append(r))
    if end is not None: end.set("sz", FS)
    pPr = p.find(qn("a:pPr"))
    for e in pPr.findall(qn("a:spcBef")): pPr.remove(e)
    sb = etree.Element(qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts"), val="300"); pPr.insert(0, sb)
    txBody.append(p)
body.top = Inches(1.4); body.left = Inches(0.4); body.width = Inches(12.5)
notes(s, "References are in IEEE style and were each checked against the publisher or proceedings page. Bracketed numbers in the Literature Review and Tools slides match this list.")

# ------------------------------------------------------------------ 19. THANK YOU
s = S[12]
tb(s, 3.0, 5.0, 7.33, 0.6, [{"runs": ["Questions & Discussion"], "align": "c"}], anchor="m", size=28, bold=True, italic=True, font=SERIF, color=MAROON)
tb(s, 1.5, 5.65, 10.33, 0.9, [{"runs": ["K. Karthikeya  ·  A. Karthik  ·  K. Likhith  ·  Y. Sandeep"], "align": "c", "after": 2},
                             {"runs": ["SENTINEL Forge  |  Text Analytics + Big Data Analytics"], "align": "c", "color": GREY}], size=17, font=SERIF)
notes(s, "Thank the audience and invite questions. Likely questions: why not fine-tune the transformer (too little data; we say so), what the gate buys you (ablation), and how honest the 17/17 is (independent review, end-to-end bridge, 11/11).")

# ------------------------------------------------------------------ ORDER
final = [S[0], S[1], S_IDEA, S[2], S[3], None, S[4], S[5], None, S[6], S[7], S[8], None, S[9], None, S[10], None, S[11], S[12]]
new = [sl for sl in prs.slides if sl not in S]   # clones in creation order: Novelty, Implementation, Results2, Timeline, Future
it = iter(new[1:])                               # skip S_IDEA (already placed)
final = [next(it) if f is None else f for f in final]
lst = prs.slides._sldIdLst
els = {int(e.get("id")): e for e in lst}
ids = [sl.slide_id for sl in final]
for e in list(lst): lst.remove(e)
for i_ in ids: lst.append(els[i_])
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
