"""Drawing helpers for the SENTINEL Forge deck (python-pptx, inches-based)."""
import copy, os
from lxml import etree
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

ICON_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "icons")

# palette: Amrita maroon dominant, navy secondary, gold accent (all from the format's own colours)
MAROON = "820019"; MAROON_D = "5E0012"; ROSE = "FBEFF1"; ROSE2 = "F3D9DE"
NAVY = "002060"; NAVY_T = "EDF1F9"
GOLD = "C99700"; GOLD_D = "8A6A00"; GOLD_T = "FFF7DB"
GREEN = "1E8E4E"; GREEN_T = "E8F5EC"; AMBER = "D98A00"; RED = "C62828"; RED_T = "FDECEC"
INK = "1F1F1F"; GREY = "595959"; LINE = "D9DCE3"; WHITE = "FFFFFF"; PAPER = "F7F8FB"

SERIF = "Times New Roman"; SANS = "Calibri"; MONO = "Consolas"

def rgb(h): return RGBColor.from_string(h)

def _shadow(shape, on=True):
    spPr = shape._element.spPr
    for e in spPr.findall(qn("a:effectLst")): spPr.remove(e)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    if on:
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="76200", dist="25400", dir="5400000", algn="t", rotWithShape="0")
        c = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(c, qn("a:alpha"), val="16000")

def box(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, shadow=False, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    _shadow(s, shadow)
    if name: s.name = name
    return s

def _run(p, txt, o):
    r = p.add_run(); r.text = txt
    f = r.font; f.size = Pt(o.get("size", 14)); f.bold = o.get("bold", False); f.italic = o.get("italic", False)
    face = o.get("font", SANS); f.name = face; f.color.rgb = rgb(o.get("color", INK))
    rPr = r._r.get_or_add_rPr()
    if o.get("spc"): rPr.set("spc", str(o["spc"]))
    cs = etree.SubElement(rPr, qn("a:cs")); cs.set("typeface", face)
    return r

def fill_tf(tf, paras, base=None):
    """paras: list of dict(runs=[str|(str,opts)], size,bold,italic,color,font,align,after,before,line,bullet)"""
    base = base or {}
    first = True
    for pd in paras:
        if isinstance(pd, str): pd = {"runs": [pd]}
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        o = {**base, **{k: v for k, v in pd.items() if k != "runs"}}
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[o.get("align", "l")]
        if "after" in o: p.space_after = Pt(o["after"])
        if "before" in o: p.space_before = Pt(o["before"])
        if "line" in o: p.line_spacing = o["line"]
        for r in pd["runs"]:
            txt, ro = (r, {}) if isinstance(r, str) else r
            _run(p, txt, {**o, **ro})
        if o.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            ind = int(o.get("indent", 0.2) * 914400)
            pPr.set("marL", str(ind)); pPr.set("indent", str(-ind))
            bc = etree.SubElement(pPr, qn("a:buClr")); etree.SubElement(bc, qn("a:srgbClr"), val=o.get("bcolor", MAROON))
            etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
            etree.SubElement(pPr, qn("a:buChar"), char=o.get("bchar", "•"))

def tb(slide, x, y, w, h, paras, anchor="t", m=(0.05, 0.03, 0.05, 0.03), **base):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(v) for v in m]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras if isinstance(paras, list) else [paras], base)
    return t

def shape_text(s, paras, anchor="m", m=(0.08, 0.04, 0.08, 0.04), **base):
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(v) for v in m]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras if isinstance(paras, list) else [paras], base)

def icon(slide, name, color, x, y, s):
    return slide.shapes.add_picture(os.path.join(ICON_DIR, f"{name}_{color}.png"), Inches(x), Inches(y), Inches(s), Inches(s))

def icon_circle(slide, name, x, y, d, bg=MAROON, fg="white", ratio=0.5):
    c = box(slide, x, y, d, d, fill=bg, shape=MSO_SHAPE.OVAL)
    k = d * ratio
    icon(slide, name, fg, x + (d - k) / 2, y + (d - k) / 2, k)
    return c

def num_circle(slide, n, x, y, d, bg=MAROON, size=16, fg=WHITE):
    c = box(slide, x, y, d, d, fill=bg, shape=MSO_SHAPE.OVAL)
    shape_text(c, [{"runs": [str(n)], "align": "c", "size": size, "bold": True, "color": fg, "font": SERIF}], m=(0, 0, 0, 0))
    return c

def chip(slide, x, y, w, h, text, fill=ROSE, color=MAROON, size=12, bold=True, line=None, font=SANS, radius=0.5):
    c = box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    shape_text(c, [{"runs": [text], "align": "c", "size": size, "bold": bold, "color": color, "font": font}], m=(0.06, 0.01, 0.06, 0.01))
    return c

def chevron(slide, x, y, w, h, color=MAROON):
    return box(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.CHEVRON)

def card(slide, x, y, w, h, fill=WHITE, line=LINE, shadow=True, radius=0.05):
    return box(slide, x, y, w, h, fill=fill, line=line, shadow=shadow, radius=radius)

def lead(slide, runs, y=1.22, size=24):
    """The format's 24 pt body line under every title."""
    if isinstance(runs, str): runs = [runs]
    return tb(slide, 0.4, y, 12.53, 0.6, [{"runs": runs}], anchor="m", size=size, color=INK)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
